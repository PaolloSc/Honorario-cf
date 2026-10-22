"""Testemunhas digitais no DOCX gerado (campos DocuSeal vs bloco fisico)."""
import re
import zipfile

from app.models.contract import ContratoRequest
from app.services.contract_generator import ContractGenerator


def _text_of(path: str) -> str:
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    paras = []
    for p in re.split(r"</w:p>", xml):
        txt = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", p))
        if txt.strip():
            paras.append(txt)
    return "\n".join(paras)


def _req() -> dict:
    return {
        "contratantes": [{
            "tipo": "PF", "nome": "Fulano", "nacionalidade": "brasileiro",
            "cpf": "00000000000", "profissao": "x", "estado_civil": "Solteiro(a)",
            "endereco": "rua x", "email": "a@a.com",
        }],
        "escopos": [{"tipo": "consultoria_lgpd", "honorarios": ["pro_labore"],
                     "pro_labore": {"valor_total": 1000, "tem_parcelamento": False}}],
        "acessorios": {"tem_reembolso": True, "tem_penalidade_inadimplemento": False},
        "participacao": {"tem_participacao": False},
    }


def test_digital_testemunha_fields_rendered():
    data = ContratoRequest(**_req())
    roles = [
        {"email": "c@a.com", "name": "Client", "role": "Contratante"},
        {"email": "lilian@cf.com", "name": "Lilian Silveira Correa", "role": "Testemunha 1"},
        {"email": "outra@cf.com", "name": "Outra Pessoa", "role": "Testemunha 2"},
    ]
    gen = ContractGenerator()
    _, path = gen.generate(data, contract_id="TESTEM_DIGITAL", signatario_roles=roles)
    text = _text_of(path)

    assert "type=signature;role=Testemunha 1" in text
    assert "type=signature;role=Testemunha 2" in text
    # Sem bloco fisico em branco (sem linhas de CPF de testemunha)
    assert "CPF:" not in text


def test_initial_generation_keeps_physical_block():
    """Sem signatario_roles -> mantem bloco fisico em branco (compat)."""
    data = ContratoRequest(**_req())
    gen = ContractGenerator()
    _, path = gen.generate(data, contract_id="TESTEM_FISICO")
    text = _text_of(path)

    assert "TESTEMUNHAS:" in text
    assert "CPF:" in text
    assert "role=Testemunha" not in text


def _linhas_da_grade(path: str) -> list[list[str]]:
    """Texto de cada celula, por linha, das tabelas do DOCX."""
    from docx import Document

    return [
        ["\n".join(p.text for p in c.paragraphs) for c in row.cells]
        for t in Document(path).tables
        for row in t.rows
    ]


def test_testemunhas_digitais_ficam_lado_a_lado():
    data = ContratoRequest(**_req())
    roles = [
        {"email": "c@a.com", "name": "Client", "role": "Contratante"},
        {"email": "lilian@cf.com", "name": "Lilian Silveira Correa", "role": "Testemunha 1"},
        {"email": "outra@cf.com", "name": "Outra Pessoa", "role": "Testemunha 2"},
    ]
    _, path = ContractGenerator().generate(data, contract_id="TESTEM_LADO", signatario_roles=roles)
    linha = next(r for r in _linhas_da_grade(path) if "role=Testemunha 1" in r[0] + r[1])
    assert "role=Testemunha 1" in linha[0] and "TESTEMUNHA 1: LILIAN" in linha[0]
    assert "role=Testemunha 2" in linha[1] and "TESTEMUNHA 2: OUTRA PESSOA" in linha[1]


def test_uma_testemunha_digital_fica_a_esquerda():
    data = ContratoRequest(**_req())
    roles = [
        {"email": "c@a.com", "name": "Client", "role": "Contratante"},
        {"email": "lilian@cf.com", "name": "Lilian Silveira Correa", "role": "Testemunha 1"},
    ]
    _, path = ContractGenerator().generate(data, contract_id="TESTEM_UMA", signatario_roles=roles)
    linha = next(r for r in _linhas_da_grade(path) if "role=Testemunha 1" in r[0] + r[1])
    assert "role=Testemunha 1" in linha[0]
    assert linha[1].strip() == ""


def test_bloco_fisico_de_testemunhas_tem_duas_colunas():
    data = ContratoRequest(**_req())
    _, path = ContractGenerator().generate(data, contract_id="TESTEM_FISICO_COLS")
    linha = next(r for r in _linhas_da_grade(path) if "CPF:" in r[0])
    assert all("Nome:" in c and "CPF:" in c for c in linha)
