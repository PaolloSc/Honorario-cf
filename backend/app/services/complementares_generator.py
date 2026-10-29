"""Geradores de aditivo, distrato e substabelecimento (issue #126).

Reaproveitam o layout e a grade de assinaturas (tags do DocuSeal) do contrato de
consumidor. O texto vem literal dos modelos do escritorio; as clausulas
condicionais e os dados variaveis estao nos metodos de cada tipo. Os dados dos
modelos de referencia (clientes, valores, processos) nao entram no codigo.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models.documentos_complementares import (
    AditivoRequest,
    DistratoRequest,
    SubstabelecimentoRequest,
    nome_contratante,
)
from app.services.consumidor_generator import ConsumidorGenerator
from app.services.contract_generator import MESES_PT_BR
from app.utils.currency import valor_com_extenso

ESCRITORIO_NOME = "CARVALHO & FURTADO ADVOGADOS"
ESCRITORIO_CNPJ = "25.463.159/0001-73"
ESCRITORIO_OAB = "OAB/MG sob o n.º 125"
ESCRITORIO_ENDERECO = "Rua Antônio de Albuquerque, n.º 271, 5º andar, Savassi, Belo Horizonte/MG, CEP 30.112-010"
LOCAL = "Belo Horizonte/MG"
# Sem signatarios definidos o documento sai com linha em branco; as tags do DocuSeal so entram no envio.
LINHA = "_" * 32


def _sem_ponto(texto: str) -> str:
    return texto.strip().rstrip(".;, ")


def _mes_ano(iso: str | None) -> str:
    """'2025-06-10' -> 'junho de 2025'."""
    d = _data(iso)
    return f"{MESES_PT_BR[d.month - 1]} de {d.year}"


def _data(iso: str | None) -> date:
    """Data ISO ja validada pelo request; vazia = hoje."""
    return date.fromisoformat(iso[:10]) if iso else date.today()


class ComplementarGenerator(ConsumidorGenerator):
    """Base: titulo centralizado, corpo do tipo, fechamento com data e assinaturas."""

    TITULO = ""

    def _build_document(self, data, signatario_roles: list[dict] | None = None) -> Document:  # type: ignore[override]
        doc = self._new_document_from_template()
        self._clear_body(doc)
        self._clear_headers_footers(doc)
        self._ensure_contract_styles(doc)

        titulo = doc.add_heading(self._titulo(data), level=1)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._corpo(doc, data)
        self._fechamento(doc, data, signatario_roles)
        self._apply_document_standard(doc)
        return doc

    # Cada tipo implementa estes tres.
    def _titulo(self, data) -> str:
        return self.TITULO

    def _corpo(self, doc: Document, data) -> None:
        raise NotImplementedError

    def _fechamento(self, doc: Document, data, signatario_roles: list[dict] | None) -> None:
        raise NotImplementedError

    # ── Partes ────────────────────────────────────────────────────

    @staticmethod
    def _cnpj_cpf(c: dict) -> str:
        if c.get("cnpj"):
            return f"inscrita no CNPJ sob o n.º {c['cnpj']}"
        return f"CPF n.º {c.get('cpf', '')}"

    def _qualificar(self, c: dict, *, com_representantes: bool = True) -> str:
        """Qualificacao do contratante a partir do form_data do contrato de origem."""
        nome = nome_contratante(c)
        if c.get("cnpj"):
            partes = [nome, "pessoa jurídica de direito privado", self._cnpj_cpf(c)]
            if c.get("endereco"):
                partes.append(f"com sede em {c['endereco']}")
            reps = [r for r in (c.get("representantes") or []) if r.get("nome")] if com_representantes else []
            if reps:
                partes.append(
                    "neste ato representada por "
                    + "; e por ".join(
                        ", ".join(filter(None, [r["nome"], f"CPF {r['cpf']}" if r.get("cpf") else None]))
                        for r in reps
                    )
                )
            return ", ".join(partes)
        partes = [nome, c.get("nacionalidade"), c.get("estado_civil"), c.get("profissao"), self._cnpj_cpf(c)]
        if c.get("endereco"):
            partes.append(f"residente em {c['endereco']}")
        return ", ".join(p for p in partes if p)

    def _data_extenso(self, iso: str | None) -> str:
        d = _data(iso)
        return self._format_date_pt_br(datetime.combine(d, datetime.min.time()))

    # ── Fechamento ────────────────────────────────────────────────

    def _data_e_local(self, doc: Document, data) -> None:
        doc.add_paragraph()
        doc.add_paragraph(f"{LOCAL}, {self._data_extenso(data.data_documento)}.")
        doc.add_paragraph()

    def _assinaturas_partes(self, doc: Document, data, signatario_roles: list[dict] | None) -> None:
        """Contratante(s) e escritorio (+ testemunhas) — aditivo e distrato."""
        entries: list[tuple[str, list[str]]] = []
        if signatario_roles:
            for sig in (s for s in signatario_roles if s.get("role", "").startswith("Contratante")):
                nome = sig.get("name", "Contratante")
                entries.append((
                    f"{{{{Assinatura {nome};type=signature;role={sig['role']}}}}}",
                    [nome.upper(), "Contratante"],
                ))
            for sig in (s for s in signatario_roles if s.get("role", "").startswith("Contratado")):
                entries.append((
                    f"{{{{Assinatura {ESCRITORIO_NOME.title()};type=signature;role={sig['role']}}}}}",
                    [ESCRITORIO_NOME, f"CNPJ n.º {ESCRITORIO_CNPJ}", "Contratada"],
                ))
        else:
            for c in data.contratantes:
                nome = nome_contratante(c)
                doc_ = c.get("cnpj") or c.get("cpf") or ""
                rotulo = [nome.upper(), f"{'CNPJ' if c.get('cnpj') else 'CPF'} n.º {doc_}", "Contratante"]
                entries.append((LINHA, rotulo))
            entries.append((
                LINHA,
                [ESCRITORIO_NOME, f"CNPJ n.º {ESCRITORIO_CNPJ}", "Contratada"],
            ))
        self._grid_assinaturas(doc, entries)

        testemunhas = [s for s in (signatario_roles or []) if s.get("role", "").startswith("Testemunha")]
        doc.add_paragraph()
        doc.add_paragraph("Testemunhas:")
        if testemunhas:
            self._grid_assinaturas(doc, [
                (
                    f"{{{{Assinatura {s.get('name', 'Testemunha')};type=signature;role={s['role']}}}}}",
                    [f"{s['role'].upper()}: {s.get('name', 'Testemunha').upper()}"],
                )
                for s in testemunhas
            ])
        else:
            doc.add_paragraph()
            doc.add_paragraph("Nome:")
            doc.add_paragraph("CPF n.º")
            doc.add_paragraph()
            doc.add_paragraph("Nome:")
            doc.add_paragraph("CPF n.º")


class AditivoGenerator(ComplementarGenerator):
    def _titulo(self, data: AditivoRequest) -> str:
        prefixo = f"{data.numero}º " if data.numero > 1 else ""
        return f"{prefixo}ADITIVO CONTRATUAL AO CONTRATO DE PRESTAÇÃO DE SERVIÇOS ADVOCATÍCIOS"

    def _corpo(self, doc: Document, data: AditivoRequest) -> None:
        doc.add_paragraph("Pelo presente instrumento particular, de um lado:")
        for c in data.contratantes:
            doc.add_paragraph(f"{self._qualificar(c)}.")
        doc.add_paragraph("E, de outro lado:")
        doc.add_paragraph(
            f"{ESCRITORIO_NOME}, pessoa jurídica de direito interno, sociedade simples de prestação de "
            f"serviços advocatícios, inscrita no CNPJ n.º {ESCRITORIO_CNPJ}, com sede na {ESCRITORIO_ENDERECO}."
        )

        # Considerandos: contrato de origem, aditivos anteriores e o atual.
        quando = f" em {_data(data.data_contrato_origem).year}" if data.data_contrato_origem else ""
        doc.add_paragraph("Considerando que:")
        doc.add_paragraph(f"As partes firmaram Contrato de Prestação de Serviços Advocatícios{quando};")
        for ant in data.aditivos_anteriores:
            doc.add_paragraph(
                f"Em {_mes_ano(ant.data_documento)}, as partes firmaram o {ant.numero}º Aditivo "
                f"Contratual, tendo por objeto {_sem_ponto(ant.objeto)};"
            )
        ordinal = f"{data.numero}º " if data.numero > 1 else ""
        doc.add_paragraph(
            f"Em {_mes_ano(data.data_documento)}, as partes acordaram o presente aditivo, "
            f"tendo por objeto {_sem_ponto(data.objeto)};"
        )
        doc.add_paragraph(
            f"Resolvem firmar o presente {ordinal}Aditivo Contratual, que se regerá pelas cláusulas e "
            "condições seguintes:"
        )

        n = 0

        def clausula(titulo: str) -> None:
            nonlocal n
            n += 1
            doc.add_heading(f"Cláusula {n} – {titulo}", level=2)

        clausula("Objeto")
        doc.add_paragraph(
            f"O presente aditivo tem por objeto a formalização de {_sem_ponto(data.objeto)}."
        )

        if data.alteracoes:
            clausula("Alterações")
            for i, alt in enumerate(data.alteracoes, 1):
                doc.add_paragraph(f"{n}.{i} – {self._redacao_alteracao(alt)}")

        clausula("Ratificação")
        doc.add_paragraph(
            "Permanecem inalteradas as demais cláusulas e condições do contrato original, que continuam "
            "em pleno vigor."
        )
        doc.add_paragraph(
            "E, por estarem assim justas e contratadas, firmam o presente aditivo em 2 (duas) vias de "
            "igual teor e forma, para um só efeito."
        )

    _MASCULINOS = {"item", "subitem", "paragrafo", "inciso", "anexo", "capitulo", "trecho"}
    _FEMININOS = {"clausula", "subclausula", "alinea", "secao", "tabela", "parte", "disposicao"}
    _ARTIGOS = {"o": ("s", "m"), "a": ("s", "f"), "os": ("p", "m"), "as": ("p", "f")}

    @classmethod
    def _dispositivo(cls, texto: str, participio: str) -> tuple[str, bool]:
        """('Fica alterada a Cláusula 3ª', False) / ('Ficam alterados os itens 3.2 e 3.3', True).

        Número e gênero vêm da primeira palavra digitada: Cláusula(s), item/itens, §, Art., ou um
        artigo já escrito ('o item 3' não leva outro artigo). Palavra desconhecida (ex.: só '3.2'):
        'a disposição'.
        """
        texto = texto.strip()
        norm = unicodedata.normalize("NFD", texto).encode("ascii", "ignore").decode().lower()
        m = re.match(r"(arts?\b|artigos?\b|[a-z]+)", norm)  # '§' some no ascii: tratado a parte
        palavra = "§§" if texto.startswith("§§") else "§" if texto.startswith("§") else (m.group(1) if m else "")
        verbo = lambda plural: "Ficam" if plural else "Fica"  # noqa: E731
        if palavra in cls._ARTIGOS:  # ja vem com artigo definido: nao duplica
            numero, genero = cls._ARTIGOS[palavra]
            artigo_minusculo = texto[0].lower() + texto[1:]  # "O item 3" -> "o item 3"
            return f"{verbo(numero == 'p')} {participio}{'o' if genero == 'm' else 'a'}{'s' if numero == 'p' else ''} {artigo_minusculo}", numero == "p"
        if palavra in ("§§", "arts", "artigos"):
            return f"Ficam {participio}os os {texto}", True
        if palavra in ("§", "art", "artigo"):
            return f"Fica {participio}o o {texto}", False
        singular = palavra
        plural = False
        if palavra.endswith("s") and palavra not in cls._MASCULINOS | cls._FEMININOS:
            for base in (palavra[:-1], palavra[:-2], palavra[:-3] + "ao" if palavra.endswith("oes") else None):
                if base in cls._MASCULINOS | cls._FEMININOS:
                    singular, plural = base, True
                    break
            if palavra in ("itens", "subitens"):
                singular, plural = palavra[:-3] + "em", True
        if singular in cls._MASCULINOS:
            return f"{verbo(plural)} {participio}o{'s' if plural else ''} {'os' if plural else 'o'} {texto}", plural
        if singular in cls._FEMININOS:
            return f"{verbo(plural)} {participio}a{'s' if plural else ''} {'as' if plural else 'a'} {texto}", plural
        return f"Fica {participio}a a disposição {texto}", False

    @classmethod
    def _redacao_alteracao(cls, alt) -> str:
        if alt.tipo == "alterar":
            frase, plural = cls._dispositivo(alt.clausula, "alterad")
            return f"{frase}, que {'passam' if plural else 'passa'} a vigorar com a seguinte redação: “{alt.texto}”"
        if alt.tipo == "excluir":
            frase, _ = cls._dispositivo(alt.clausula, "excluíd")
            return f"{frase} do contrato original."
        return f"Fica incluída no contrato original a seguinte cláusula: “{alt.texto}”"

    def _fechamento(self, doc: Document, data: AditivoRequest, signatario_roles) -> None:
        self._data_e_local(doc, data)
        self._assinaturas_partes(doc, data, signatario_roles)


class DistratoGenerator(ComplementarGenerator):
    """Texto do modelo do escritorio para distrato; os dados do modelo de referencia nao entram no codigo."""

    TITULO = "DISTRATO DE PRESTAÇÃO DE SERVIÇOS ADVOCATÍCIOS"

    def _corpo(self, doc: Document, data: DistratoRequest) -> None:
        for c in data.contratantes:
            doc.add_paragraph(f"{self._qualificar(c, com_representantes=False)}; denominado apenas Contratante;")
        doc.add_paragraph(
            f"{ESCRITORIO_NOME}, pessoa jurídica de direito interno, sociedade simples de prestação de "
            f"serviços advocatícios, inscrita no CNPJ n.º {ESCRITORIO_CNPJ}, com sede na "
            f"{ESCRITORIO_ENDERECO}, doravante denominado apenas Contratado."
        )
        doc.add_paragraph("Têm entre si contratadas as seguintes cláusulas do presente Distrato:")

        n = 0

        def clausula(texto: str) -> int:
            nonlocal n
            n += 1
            doc.add_paragraph(f"{n}. {texto}")
            return n

        clausula(
            f"As partes declaram e reconhecem que firmaram, entre si, em {self._data_extenso(data.data_contrato_origem)}, "
            "Contrato de Prestação de Serviço Advocatício para "
            f"{_sem_ponto(data.objeto_contrato_origem)}."
        )
        for ant in data.aditivos_anteriores:
            doc.add_paragraph(
                f"Reconhecem, ainda, o {ant.numero}º Aditivo Contratual ao referido Contrato, firmado em "
                f"{_mes_ano(ant.data_documento)}, tendo por objeto {_sem_ponto(ant.objeto)}."
            )
        k = clausula(
            "Por motivos supervenientes à assinatura do referido instrumento, as partes decidem de comum "
            "acordo firmar o presente Distrato nas seguintes condições:"
        )
        sub = 0
        if data.pagamento_valor:
            sub += 1
            forma = _sem_ponto(data.pagamento_forma) or f"Pix (CNPJ {ESCRITORIO_CNPJ})"
            doc.add_paragraph(
                f"{k}.{sub}. A Contratante pagará ao Contratado, até o dia {self._data_curta(data.pagamento_data)}, "
                f"o valor de {valor_com_extenso(data.pagamento_valor)}, por meio de {forma}."
            )
        if data.substabelecer_e_devolver:
            sub += 1
            doc.add_paragraph(
                f"{k}.{sub}. O Contratado compromete-se a substabelecer sem reserva de poderes a favor do novo "
                "patrono a ser indicado pela Contratante, bem como a restituir eventuais documentos físicos ou "
                "digitais que estejam sob sua guarda."
            )

        # Com ressalva (opcao do formulario), a quitacao exclui os honorarios e reembolsos vincendos.
        ressalva = (
            ", ressalvados reembolsos e honorários vincendos, inclusive de êxito, caso exigíveis nos termos do contrato original"
            if data.ressalva_exito else ""
        )
        clausula(
            "Salvo as obrigações previstas neste Distrato, as partes, com o presente instrumento, se dão mútua, "
            "total, completa, irrevogável e irretratável quitação, declarando que nada têm a reclamar extra ou "
            f"judicialmente uma da outra, renunciando a qualquer direito ou pretensão nesse sentido{ressalva}."
        )
        clausula(
            "O Contrato objeto do presente Distrato resta extinto, sem qualquer ônus, penalidade, dever e/ou "
            "obrigação adicionais entre si, além dos expressamente pactuados neste instrumento, reconhecendo "
            "ambas as Partes que não há pendências, créditos ou débitos entre elas além dos aqui previstos."
        )
        clausula(
            "A tolerância de qualquer das partes com descumprimento de obrigações constantes no presente "
            "instrumento não deverá ser tido como novação ou ter como efeito a mudança de qualquer dos termos "
            "acordados no presente Distrato."
        )
        clausula(
            "O presente instrumento incorpora todas as tratativas, acordos e combinados entre as partes, "
            "excluindo qualquer outra que nele não conste expressamente."
        )
        clausula("O foro da comarca eleita no contrato permanece o competente para dirimir qualquer conflito oriundo deste Distrato.")
        doc.add_paragraph(
            "E, por estarem as partes justas e contratadas e por considerarem o negócio perfeito, assinam o "
            "presente instrumento em 02 (duas) vias de igual teor e forma, para um só efeito, juntamente com "
            "as testemunhas abaixo."
        )

    @staticmethod
    def _data_curta(iso: str | None) -> str:
        return _data(iso).strftime("%d/%m/%Y")

    def _fechamento(self, doc: Document, data: DistratoRequest, signatario_roles) -> None:
        self._data_e_local(doc, data)
        self._assinaturas_partes(doc, data, signatario_roles)


class SubstabelecimentoGenerator(ComplementarGenerator):
    TITULO = "SUBSTABELECIMENTO SEM RESERVA DE PODERES"

    @staticmethod
    def _qualificar_substabelecido(s) -> str:
        """Sem genero (formulario atual) o texto sai neutro; com genero (legado), como antes."""
        if s.genero is None:
            return f"{s.nome}, inscrito(a) na {s.oab}"
        f = s.genero == "F"
        nacionalidade = s.nacionalidade.strip() or ("brasileira" if f else "brasileiro")
        partes = [f"{'Dra.' if f else 'Dr.'} {s.nome}", nacionalidade]
        if s.estado_civil.strip():
            partes.append(s.estado_civil.strip())
        partes.append(f"{'inscrita' if f else 'inscrito'} na {s.oab}")
        return ", ".join(partes)

    def _corpo(self, doc: Document, data: SubstabelecimentoRequest) -> None:
        substabelecentes = "; ".join(f"{s.nome.upper()}, {s.oab}" for s in data.substabelecentes)
        substabelecidos = "; ".join(self._qualificar_substabelecido(s) for s in data.substabelecidos)
        verbo = "substabelece" if len(data.substabelecentes) == 1 else "substabelecem"
        sem_genero = all(s.genero is None for s in data.substabelecidos)
        if len(data.substabelecidos) == 1:
            if sem_genero:
                artigo = "ao(à) advogado(a)"
            else:
                artigo = "à advogada" if data.substabelecidos[0].genero == "F" else "ao advogado"
            endereco = "com endereço profissional à"
        else:
            artigo, endereco = "aos advogados", "todos com endereço profissional à"
        local = _sem_ponto(data.endereco_substabelecidos)
        # Endereco e' opcional: sem ele a oracao segue direto para os poderes.
        onde = f", {endereco} {local}," if local else ","
        doc.add_paragraph(
            f"Pelo presente instrumento, {substabelecentes}, {verbo}, SEM RESERVA DE PODERES, {artigo} "
            f"{substabelecidos}{onde} os poderes que {'lhe foram conferidos' if len(data.substabelecentes) == 1 else 'lhes foram conferidos'} por "
            f"{data.outorgante.exibicao}, no âmbito dos processos judiciais abaixo indicados, seus "
            "incidentes, desdobramentos e recursos correlatos, ressalvados os honorários contratuais e "
            "sucumbenciais:"
        )
        for p in data.processos:
            sufixo = f" – {_sem_ponto(p.descricao)}" if p.descricao.strip() else ""
            doc.add_paragraph(f"{p.tipo} n.º {p.numero}{sufixo}.")
        doc.add_paragraph(
            "Os efeitos do presente substabelecimento se estendem a todos os patronos cujos poderes foram "
            "outorgados à época que possuíam vínculo com o escritório Carvalho & Furtado Advogados, "
            f"inscrita no CNPJ n.º {ESCRITORIO_CNPJ} e na {ESCRITORIO_OAB}, com sede na Rua Antônio de "
            "Albuquerque, 271, 5º andar, Savassi, Belo Horizonte/MG, CEP 30.112-010."
        )

    def _fechamento(self, doc: Document, data: SubstabelecimentoRequest, signatario_roles) -> None:
        self._data_e_local(doc, data)
        if signatario_roles:
            sigs = [s for s in signatario_roles if s.get("role", "").startswith("Substabelecente")]
            por_email = {s.email.lower(): s for s in data.substabelecentes}
            entries = []
            for sig in sigs:
                nome = sig.get("name", "Substabelecente")
                cadastro = por_email.get((sig.get("email") or "").lower())
                rotulo = [nome.upper()] + ([cadastro.oab] if cadastro else []) + ["Substabelecente"]
                entries.append((f"{{{{Assinatura {nome};type=signature;role={sig['role']}}}}}", rotulo))
        else:
            entries = [(LINHA, [s.nome.upper(), s.oab, "Substabelecente"]) for s in data.substabelecentes]
        self._grid_assinaturas(doc, entries)


def _romano(n: int) -> str:
    out = ""
    for valor, letra in ((10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
        while n >= valor:
            out += letra
            n -= valor
    return out
