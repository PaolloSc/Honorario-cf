"""Aceitacao da issue #126: aditivo, distrato e substabelecimento na plataforma.

Testes escritos ANTES da implementacao — todos falham ate a feature existir.
Um seam so: a API HTTP. Gera o documento, baixa o .docx e confere o texto;
envia para assinatura com o DocuSeal mockado; simula o webhook do distrato.

Contrato de API fixado por estes testes:
  POST /api/documentos/{tipo}              tipo in aditivo|distrato|substabelecimento
       -> {success, contract_id, download_url}
  GET  /api/documentos/prefill/{contrato_origem_id}
       -> {contratantes, aditivos_existentes}
  GET  /api/contracts/{contract_id}/documentos
       -> {documentos: [{contract_id, tipo_contrato, status}]}

Formato fixado: dia com 2 digitos ("05 de outubro de 2026"), como nos modelos.
OAB vai em texto livre e e impressa como veio. Numero de processo e CNJ com
digito verificador; "tipo" do processo e livre (Processo, Agravo de Instrumento...).

Todos os dados sao ficticios (processos com DV valido inventados).
"""
import io
import json
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document
from fastapi import Request

from app.auth import CurrentUser, _enforce_readonly, get_current_user
from app.database import ColaboradorDB, ContractDB, ContractVersionDB, SessionLocal, utcnow
from app.main import app

SECRET = "test-webhook-secret"
ORIGEM = "origem-001"
SOCIO = {"email": "socia@cf.com", "name": "Sócia CF", "role": "Contratado"}
CONTRATANTE_SIG = {"email": "c@exemplo.com", "name": "Fulano de Tal", "role": "Contratante"}

CLIENTE_PJ = {
    "tipo": "PJ",
    "razao_social": "EMPRESA EXEMPLO LTDA",
    "cnpj": "11.222.333/0001-81",
    "endereco": "Rua Exemplo, 100, Centro, Contagem/MG, CEP 32.000-000",
    "email": "contato@exemplo.com",
    "representantes": [{"nome": "Fulano de Tal", "cpf": "123.456.789-09"}],
}

CLIENTE_PF = {
    "tipo": "PF",
    "nome": "Mônica Exemplo da Silva",
    "nacionalidade": "brasileira",
    "cpf": "987.654.321-00",
    "profissao": "engenheira",
    "estado_civil": "casado",
    "endereco": "Rua das Acácias, 50, Belo Horizonte/MG, CEP 30.100-000",
    "email": "monica@exemplo.com",
}

ADITIVO = {
    "contrato_origem_id": ORIGEM,
    "data_documento": "2026-10-05",
    "data_contrato_origem": "2020-01-15",
    "objeto": "prestação de serviços de auditoria trabalhista",
    "alteracoes": [
        {"tipo": "alterar", "clausula": "Cláusula 3ª, item 3.2", "texto": "O valor mensal passa a ser de R$ 2.000,00."},
        {"tipo": "excluir", "clausula": "Cláusula 5ª"},
        {"tipo": "adicionar", "texto": "As partes poderão ajustar metas por escrito."},
    ],
}

# Aditivo salvo antes de 22/10/2026 (remuneracao/abrangencia): precisa continuar abrindo e gerando.
ADITIVO_LEGADO = {
    **{k: v for k, v in ADITIVO.items() if k != "alteracoes"},
    "remuneracao": {"valor_mensal": 1250.5, "franquia_horas": 1, "valor_hora_excedente": 210.0},
    "compensacao_anual": True,
    "abrangencia": [{"razao_social": "FILIAL ÁGUA BRANCA LTDA", "cnpj": "44.555.666/0001-77"}],
}

DISTRATO = {
    "contrato_origem_id": ORIGEM,
    "data_documento": "2026-10-07",
    "data_contrato_origem": "2021-03-09",
    "objeto_contrato_origem": "assessoria jurídica empresarial",
    "ressalva_exito": False,
    "substabelecer_e_devolver": True,
    "substabelecimento": {
        "substabelecentes": [{"nome": "Sócia CF", "oab": "OAB/MG n.º 111.111", "email": "socia@cf.com"}],
        "substabelecidos": [{"nome": "Novo Patrono", "oab": "OAB/SP sob o n.º 55.555"}],
        "endereco_substabelecidos": "",
        "processos": [{"tipo": "Processo", "numero": "1234567-49.2020.8.13.0024", "descricao": "Cliente x Parte"}],
    },
}

PROC_A = {"tipo": "Processo", "numero": "1234567-49.2020.8.13.0024",
          "descricao": "Outorgante x Parte Contrária A"}
PROC_B = {"tipo": "Agravo de Instrumento", "numero": "1111111-60.2023.4.06.0000",
          "descricao": "Outorgante x Parte Contrária B"}

SUBSTABELECIMENTO = {
    "contrato_origem_id": None,
    "data_documento": "2026-10-08",
    "outorgante": {"razao_social": "OUTORGANTE EXEMPLO LTDA"},
    "substabelecentes": [
        {"nome": "Sócia CF", "oab": "OAB/MG nº 111.111", "email": "socia@cf.com"},
        {"nome": "Advogado CF", "oab": "OAB/MG nº 222.222", "email": "adv@cf.com"},
    ],
    "substabelecidos": [
        {"nome": "Beltrana Parceira", "genero": "F", "estado_civil": "casada",
         "oab": "OAB/MG sob o nº 333.333"},
        {"nome": "Ciclano Parceiro", "genero": "M", "estado_civil": "solteiro",
         "oab": "OAB/MG sob o nº 444.444"},
    ],
    "endereco_substabelecidos": "Rua Parceira, 10, Belo Horizonte/MG, CEP 30.000-000",
    "processos": [PROC_A, PROC_B],
}


# ── Helpers ───────────────────────────────────────────────────────────────


def _user(role: str = "advogado", email: str = "lawyer@test.com"):
    return lambda: CurrentUser(azure_id="x", email=email, name="Lawyer", role=role)


def _semear_origem(contratante: dict, status: str = "assinado"):
    db = SessionLocal()
    try:
        db.query(ContractVersionDB).filter_by(contract_id=ORIGEM).delete()
        db.query(ContractDB).filter_by(contract_id=ORIGEM).delete()
        db.commit()
        db.add(ContractDB(
            contract_id=ORIGEM, status=status, client_name="Cliente",
            client_email=contratante["email"], current_version=1, created_by="lawyer@test.com",
            created_at=utcnow(), updated_at=utcnow(),
        ))
        db.commit()
        db.add(ContractVersionDB(
            contract_id=ORIGEM, version_number=1,
            form_data_json=json.dumps({"contratantes": [contratante], "escopos": []}),
            created_at=utcnow(),
        ))
        db.commit()
    finally:
        db.close()


@pytest.fixture(autouse=True)
def auth_e_origem():
    app.dependency_overrides[get_current_user] = _user()
    db = SessionLocal()
    try:
        db.add(ColaboradorDB(nome=SOCIO["name"], email=SOCIO["email"], papel="socio", ativo=True))
        db.commit()
    finally:
        db.close()
    _semear_origem(CLIENTE_PJ)
    yield
    app.dependency_overrides.pop(get_current_user, None)


def _gerar(client, tipo: str, payload: dict):
    return client.post(f"/api/documentos/{tipo}", json=payload)


def _id(client, tipo: str, payload: dict) -> str:
    r = _gerar(client, tipo, payload)
    assert r.status_code == 200, r.text
    return r.json()["contract_id"]


def _texto(client, contract_id: str) -> str:
    r = client.get(f"/api/contract/{contract_id}/download")
    assert r.status_code == 200, r.text
    doc = Document(io.BytesIO(r.content))
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for linha in t.rows:
            partes.extend(c.text for c in linha.cells)
    texto = "\n".join(partes)
    _sem_marcadores(texto)
    return texto


def _sem_marcadores(texto: str):
    """Nenhum documento sai com campo nao preenchido (o modelo real tem '[número de vias]')."""
    for lixo in ("[", "{{", "None", "R$ 0,00", "0 (zero)"):
        assert lixo not in texto, f"marcador solto: {lixo!r}"


def _status(contract_id: str) -> str:
    db = SessionLocal()
    try:
        return db.query(ContractDB).filter_by(contract_id=contract_id).one().status
    finally:
        db.close()


def _enviar_capturando(client, contract_id: str, signatarios: list[dict]) -> list[dict]:
    capturados: list[dict] = []

    async def capture(template_id, signatarios, send_email=True):
        capturados.extend(signatarios)
        return {"success": True, "submission": {"id": 999, "submitters": []}, "message": "ok"}

    svc = MagicMock()
    svc.create_template_from_docx = AsyncMock(return_value={"id": 1, "name": "T"})
    svc.send_for_signature = AsyncMock(side_effect=capture)
    with patch("app.routers.docuseal.get_docuseal_service", return_value=svc), \
         patch("app.routers.docuseal._send_contract_to_financeiro", new_callable=AsyncMock), \
         patch("app.routers.docuseal._send_participacao_to_financeiro", new_callable=AsyncMock):
        r = client.post("/api/docuseal/send-for-signature",
                        json={"contract_id": contract_id, "signatarios": signatarios})
    assert r.status_code == 200, r.text
    return capturados


def _webhook_concluido(client):
    with patch("app.routers.docuseal.DOCUSEAL_WEBHOOK_SECRET", SECRET), \
         patch("app.routers.docuseal._send_participacao_to_financeiro", new_callable=AsyncMock):
        r = client.post("/api/docuseal/webhook",
                        json={"event_type": "submission.completed", "data": {"id": "999"}},
                        headers={"x-docuseal-secret": SECRET})
    assert r.status_code == 200


# ── Pre-preenchimento ─────────────────────────────────────────────────────


def test_prefill_traz_partes_do_contrato_de_origem(client):
    r = client.get(f"/api/documentos/prefill/{ORIGEM}")
    assert r.status_code == 200
    body = r.json()
    assert body["contratantes"][0]["cnpj"] == CLIENTE_PJ["cnpj"]
    assert body["aditivos_existentes"] == 0


# ── Aditivo ───────────────────────────────────────────────────────────────


def test_aditivo_gera_docx_com_partes_e_alteracoes_dos_tres_tipos(client):
    texto = _texto(client, _id(client, "aditivo", ADITIVO))

    assert "ADITIVO CONTRATUAL AO CONTRATO DE PRESTAÇÃO DE SERVIÇOS ADVOCATÍCIOS" in texto
    assert "º ADITIVO" not in texto  # o 1º aditivo nao leva numero no titulo
    assert "EMPRESA EXEMPLO LTDA" in texto and "11.222.333/0001-81" in texto
    assert "CARVALHO & FURTADO ADVOGADOS" in texto
    assert "prestação de serviços de auditoria trabalhista" in texto
    assert ("2.1 – Fica alterada a Cláusula 3ª, item 3.2, que passa a vigorar com a seguinte redação: "
            "“O valor mensal passa a ser de R$ 2.000,00.”") in texto
    assert "2.2 – Fica excluída a Cláusula 5ª do contrato original." in texto
    assert ("2.3 – Fica incluída no contrato original a seguinte cláusula: "
            "“As partes poderão ajustar metas por escrito.”") in texto
    assert "Permanecem inalteradas as demais cláusulas" in texto
    assert "vias de igual teor" in texto
    assert "05 de outubro de 2026" in texto
    # secoes removidas
    assert "Remuneração" not in texto and "Abrangência" not in texto and "Compensação" not in texto


def test_aditivo_varias_alteracoes_do_mesmo_tipo_ficam_numeradas(client):
    payload = {**ADITIVO, "alteracoes": [
        {"tipo": "excluir", "clausula": "Cláusula 4ª"},
        {"tipo": "excluir", "clausula": "Cláusula 6ª, item 6.1"},
    ]}
    texto = _texto(client, _id(client, "aditivo", payload))
    assert "2.1 – Fica excluída a Cláusula 4ª do contrato original." in texto
    assert "2.2 – Fica excluída a Cláusula 6ª, item 6.1 do contrato original." in texto


def test_aditivo_legado_com_remuneracao_ainda_gera_sem_as_secoes_antigas(client):
    texto = _texto(client, _id(client, "aditivo", ADITIVO_LEGADO))
    assert "prestação de serviços de auditoria trabalhista" in texto
    assert "Remuneração" not in texto and "FILIAL ÁGUA BRANCA" not in texto
    assert "Permanecem inalteradas as demais cláusulas" in texto


@pytest.mark.parametrize("alteracao", [
    {"tipo": "alterar", "clausula": "Cláusula 3ª"},  # sem novo texto
    {"tipo": "alterar", "texto": "novo"},  # sem clausula
    {"tipo": "excluir"},
    {"tipo": "adicionar", "texto": "  "},
    {"tipo": "inventar", "texto": "x"},
])
def test_aditivo_alteracao_invalida_e_rejeitada(client, alteracao):
    assert _gerar(client, "aditivo", {**ADITIVO, "alteracoes": [alteracao]}).status_code == 422


def test_aditivo_novo_exige_ao_menos_uma_alteracao(client):
    assert _gerar(client, "aditivo", {**ADITIVO, "alteracoes": []}).status_code == 422


def test_aditivo_pf_qualifica_com_cpf(client):
    _semear_origem(CLIENTE_PF)
    texto = _texto(client, _id(client, "aditivo", ADITIVO))
    assert "Mônica Exemplo da Silva" in texto and "987.654.321-00" in texto


def _assinar(contract_id: str):
    """Atalho: o fluxo de assinatura em si e' coberto pelos testes de webhook."""
    db = SessionLocal()
    try:
        db.query(ContractDB).filter_by(contract_id=contract_id).update({"status": "assinado"})
        db.commit()
    finally:
        db.close()


def test_aditivos_seguintes_sao_numerados_e_citam_os_anteriores(client):
    _assinar(_id(client, "aditivo", ADITIVO))
    segundo_id = _id(client, "aditivo", {**ADITIVO, "objeto": "elaboração de políticas internas"})
    segundo = _texto(client, segundo_id)
    _assinar(segundo_id)
    terceiro = _texto(client, _id(client, "aditivo", {**ADITIVO, "objeto": "treinamento de gestores"}))

    assert "2º ADITIVO CONTRATUAL" in segundo
    assert "auditoria trabalhista" in segundo  # considerando do 1º
    assert "3º ADITIVO CONTRATUAL" in terceiro
    assert "auditoria trabalhista" in terceiro and "políticas internas" in terceiro
    assert client.get(f"/api/documentos/prefill/{ORIGEM}").json()["aditivos_existentes"] == 3


def test_aditivo_exige_contrato_de_origem(client):
    assert _gerar(client, "aditivo", {**ADITIVO, "contrato_origem_id": None}).status_code == 422


def test_aditivo_com_origem_inexistente_da_404(client):
    r = _gerar(client, "aditivo", {**ADITIVO, "contrato_origem_id": "nao-existe"})
    assert r.status_code == 404
    assert r.json()["detail"] != "Not Found"  # 404 do contrato, nao de rota inexistente


def test_aditivo_e_assinado_por_contratante_socio_e_testemunhas(client):
    cid = _id(client, "aditivo", ADITIVO)
    papeis = [s["role"] for s in _enviar_capturando(client, cid, [CONTRATANTE_SIG, SOCIO])]
    assert any(p.startswith("Contratante") for p in papeis)
    assert "Contratado" in papeis
    assert any(p.startswith("Testemunha") for p in papeis)


def test_aditivo_assinado_nao_muda_status_da_origem(client):
    cid = _id(client, "aditivo", ADITIVO)
    _enviar_capturando(client, cid, [CONTRATANTE_SIG, SOCIO])
    _webhook_concluido(client)
    assert _status(cid) == "assinado"
    assert _status(ORIGEM) == "assinado"


# ── Distrato ──────────────────────────────────────────────────────────────


def test_distrato_quitacao_total_com_substabelecer(client):
    texto = _texto(client, _id(client, "distrato", DISTRATO))

    assert "DISTRATO DE PRESTAÇÃO DE SERVIÇOS ADVOCATÍCIOS" in texto
    assert "09 de março de 2021" in texto
    assert "assessoria jurídica empresarial" in texto
    assert "Por motivos supervenientes" in texto
    assert RESSALVA not in texto  # quitacao total
    assert "substabelecerá, sem reserva de poderes" in texto
    assert "02 (duas) vias" in texto and "testemunhas abaixo" in texto
    assert "TESTEMUNHAS" in texto.upper()
    assert "07 de outubro de 2026" in texto
    assert "Aditivo" not in texto  # sem aditivos cadastrados, nao cita nenhum


# Redacao aprovada pelo socio (HCF-126).
RESSALVA = ("ressalvados os honorários de êxito, que serão devidos de forma proporcional ao trabalho realizado até a data "
            "deste Distrato")
@pytest.mark.parametrize("ressalva,substabelecer", [
    (True, True), (True, False), (False, False),
])
def test_distrato_flags_sao_independentes(client, ressalva, substabelecer):
    payload = {**DISTRATO, "ressalva_exito": ressalva, "substabelecer_e_devolver": substabelecer}
    texto = _texto(client, _id(client, "distrato", payload))

    assert (RESSALVA in texto) is ressalva
    assert "mútua, total, completa, irrevogável e irretratável quitação" in texto  # com e sem ressalva
    assert ("substabelecerá, sem reserva de poderes" in texto) is substabelecer


def test_distrato_cita_so_aditivos_assinados(client):
    # Decisao: aditivo em rascunho conta na numeracao, mas o distrato so cita
    # como "celebrado" o que foi assinado.
    assinado = _id(client, "aditivo", ADITIVO)
    _enviar_capturando(client, assinado, [CONTRATANTE_SIG, SOCIO])
    _webhook_concluido(client)
    _id(client, "aditivo", {**ADITIVO, "objeto": "treinamento de gestores"})  # fica em rascunho

    texto = _texto(client, _id(client, "distrato", DISTRATO))
    assert "auditoria trabalhista" in texto
    assert "treinamento de gestores" not in texto and "2º Aditivo" not in texto


def test_distrato_duplicado_e_bloqueado(client):
    assert _gerar(client, "distrato", DISTRATO).status_code == 200
    assert _gerar(client, "distrato", DISTRATO).status_code == 409


def test_distrato_e_assinado_por_contratante_socio_e_testemunhas(client):
    cid = _id(client, "distrato", DISTRATO)
    papeis = [s["role"] for s in _enviar_capturando(client, cid, [CONTRATANTE_SIG, SOCIO])]
    assert any(p.startswith("Contratante") for p in papeis)
    assert "Contratado" in papeis
    assert any(p.startswith("Testemunha") for p in papeis)


def test_distrato_enviado_mas_nao_assinado_nao_encerra_origem(client):
    cid = _id(client, "distrato", DISTRATO)
    _enviar_capturando(client, cid, [CONTRATANTE_SIG, SOCIO])
    assert _status(ORIGEM) == "assinado"


def test_distrato_assinado_encerra_o_contrato_de_origem(client):
    cid = _id(client, "distrato", DISTRATO)
    _enviar_capturando(client, cid, [CONTRATANTE_SIG, SOCIO])
    _webhook_concluido(client)
    assert _status(cid) == "assinado"
    assert _status(ORIGEM) == "encerrado"


# ── Substabelecimento ─────────────────────────────────────────────────────


def test_substabelecimento_sem_origem_lista_processos_e_advogados(client):
    texto = _texto(client, _id(client, "substabelecimento", SUBSTABELECIMENTO))

    assert "SUBSTABELECIMENTO SEM RESERVA DE PODERES" in texto
    assert "SÓCIA CF, OAB/MG nº 111.111" in texto  # substabelecentes em MAIUSCULAS
    assert "ADVOGADO CF, OAB/MG nº 222.222" in texto
    assert "substabelecem, SEM RESERVA DE PODERES" in texto
    assert "Dra. Beltrana Parceira, brasileira, casada, inscrita na OAB/MG sob o nº 333.333" in texto
    assert "Dr. Ciclano Parceiro, brasileiro, solteiro, inscrito na OAB/MG sob o nº 444.444" in texto
    assert "Rua Parceira, 10, Belo Horizonte/MG, CEP 30.000-000" in texto
    assert "OUTORGANTE EXEMPLO LTDA" in texto
    assert "seus incidentes, desdobramentos e recursos correlatos" in texto
    assert "ressalvados os honorários contratuais e sucumbenciais" in texto
    assert "Processo n.º 1234567-49.2020.8.13.0024" in texto and "Parte Contrária A" in texto
    assert "Agravo de Instrumento n.º 1111111-60.2023.4.06.0000" in texto and "Parte Contrária B" in texto
    assert "25.463.159/0001-73" in texto  # extensao fixa a patronos anteriores
    assert "08 de outubro de 2026" in texto


def test_substabelecimento_com_um_processo(client):
    texto = _texto(client, _id(client, "substabelecimento", {**SUBSTABELECIMENTO, "processos": [PROC_A]}))
    assert "1234567-49.2020.8.13.0024" in texto and "1111111-60.2023.4.06.0000" not in texto


@pytest.mark.parametrize("numero", [
    "123",                          # formato invalido
    "1234567-50.2020.8.13.0024",    # digito verificador errado
])
def test_substabelecimento_rejeita_numero_de_processo_invalido(client, numero):
    payload = {**SUBSTABELECIMENTO, "processos": [{**PROC_A, "numero": numero}]}
    assert _gerar(client, "substabelecimento", payload).status_code == 422


@pytest.mark.parametrize("campo", ["processos", "substabelecentes", "substabelecidos"])
def test_substabelecimento_exige_listas_nao_vazias(client, campo):
    assert _gerar(client, "substabelecimento", {**SUBSTABELECIMENTO, campo: []}).status_code == 422


def test_substabelecimento_e_assinado_so_pelos_substabelecentes(client):
    sid = _id(client, "substabelecimento", SUBSTABELECIMENTO)
    enviados = _enviar_capturando(client, sid, [
        {"email": "socia@cf.com", "name": "Sócia CF", "role": "Substabelecente"},
        {"email": "adv@cf.com", "name": "Advogado CF", "role": "Substabelecente"},
    ])
    papeis = [s["role"] for s in enviados]
    assert not any(p.startswith("Testemunha") for p in papeis)
    assert not any(p.startswith("Contratante") for p in papeis)
    assert {s["email"] for s in enviados} == {"socia@cf.com", "adv@cf.com"}


# ── Vinculo, listagem e acesso ────────────────────────────────────────────


def test_documentos_derivados_aparecem_no_contrato_de_origem(client):
    _id(client, "aditivo", ADITIVO)
    _id(client, "distrato", DISTRATO)
    r = client.get(f"/api/contracts/{ORIGEM}/documentos")
    assert r.status_code == 200
    tipos = sorted(d["tipo_contrato"] for d in r.json()["documentos"])
    assert tipos == ["aditivo", "distrato", "substabelecimento"]  # o distrato gera o substabelecimento junto


def test_edicao_de_rascunho_regenera_nova_versao(client):
    # Amarra o formato salvo: form_data usa as mesmas chaves do payload.
    cid = _id(client, "aditivo", ADITIVO)
    form = client.get(f"/api/contracts/{cid}/form-data").json()["form_data"]
    form["objeto"] = "objeto corrigido"
    assert client.put(f"/api/contracts/{cid}", json={"form_data": form}).status_code == 200
    assert "objeto corrigido" in _texto(client, cid)
    assert client.get(f"/api/contracts/{cid}").json()["current_version"] == 2


def _origem_de_outro_advogado():
    db = SessionLocal()
    try:
        db.query(ContractDB).filter_by(contract_id=ORIGEM).update({"created_by": "outro@test.com"})
        db.commit()
    finally:
        db.close()


def test_advogado_nao_acessa_contrato_alheio(client):
    _origem_de_outro_advogado()
    assert _gerar(client, "aditivo", ADITIVO).status_code == 403
    assert client.get(f"/api/documentos/prefill/{ORIGEM}").status_code == 403  # expoe CPF/CNPJ
    assert client.get(f"/api/contracts/{ORIGEM}/documentos").status_code == 403


def test_admin_acessa_contrato_alheio(client):
    _origem_de_outro_advogado()
    app.dependency_overrides[get_current_user] = _user("admin", "admin@test.com")
    assert client.get(f"/api/documentos/prefill/{ORIGEM}").status_code == 200
    assert _gerar(client, "aditivo", ADITIVO).status_code == 200


def test_leitor_nao_cria_documento(client):
    # O override pula o guard de somente-leitura; aqui ele e aplicado de verdade,
    # provando que o router novo depende de get_current_user.
    def leitor(request: Request):
        return _enforce_readonly(request, CurrentUser(azure_id="x", email="l@test.com",
                                                      name="Leitor", role="leitor"))

    app.dependency_overrides[get_current_user] = leitor
    assert _gerar(client, "aditivo", ADITIVO).status_code == 403


def test_tipo_desconhecido_da_404(client):
    r = _gerar(client, "procuracao", ADITIVO)
    assert r.status_code == 404
    assert "procuracao" in r.json()["detail"]  # 404 do tipo, nao de rota inexistente


# ── Ajustes da revisao do PR #127 ─────────────────────────────────────────


def test_distrato_com_pessoa_fisica_nao_fala_em_contrato_social(client):
    _semear_origem(CLIENTE_PF)
    texto = _texto(client, _id(client, "distrato", DISTRATO))
    assert "denominado(a) apenas CONTRATANTE" in texto
    assert "inscrito(a)" not in texto
    assert "Contrato Social" not in texto  # pessoa fisica nao tem; o modelo nao cita o do escritorio


def _data_do_contrato_na_origem(iso: str):
    db = SessionLocal()
    try:
        ver = db.query(ContractVersionDB).filter_by(contract_id=ORIGEM).one()
        form = json.loads(ver.form_data_json)
        form["data_contrato_origem"] = iso
        ver.form_data_json = json.dumps(form)
        db.commit()
    finally:
        db.close()


# Decisao do usuario: data do contrato de origem e' OBRIGATORIA no aditivo e no
# distrato, para nao sair considerando sem ano nem com ano errado.
@pytest.mark.parametrize("tipo,payload", [("aditivo", ADITIVO), ("distrato", DISTRATO)])
def test_data_do_contrato_de_origem_e_obrigatoria(client, tipo, payload):
    sem = {k: v for k, v in payload.items() if k != "data_contrato_origem"}
    assert _gerar(client, tipo, sem).status_code == 422
    assert _gerar(client, tipo, {**payload, "data_contrato_origem": None}).status_code == 422


def test_aditivo_cita_o_ano_informado(client):
    assert "Serviços Advocatícios em 2020;" in _texto(client, _id(client, "aditivo", ADITIVO))


def test_prefill_sugere_a_data_do_contrato_do_pai(client):
    assert client.get(f"/api/documentos/prefill/{ORIGEM}").json()["data_contrato_origem"] is None
    _data_do_contrato_na_origem("2019-05-02")
    assert client.get(f"/api/documentos/prefill/{ORIGEM}").json()["data_contrato_origem"] == "2019-05-02"


def test_aditivo_prefere_a_data_informada_pelo_advogado(client):
    _data_do_contrato_na_origem("2019-05-02")
    texto = _texto(client, _id(client, "aditivo", ADITIVO))
    assert "Serviços Advocatícios em 2020;" in texto and "2019" not in texto


@pytest.mark.parametrize("tipo,payload,campo", [
    ("aditivo", ADITIVO, "data_documento"),
    ("distrato", DISTRATO, "data_documento"),
    ("distrato", DISTRATO, "data_contrato_origem"),
    ("aditivo", ADITIVO, "data_contrato_origem"),
    ("substabelecimento", SUBSTABELECIMENTO, "data_documento"),
])
def test_data_invalida_da_422(client, tipo, payload, campo):
    assert _gerar(client, tipo, {**payload, campo: "2026-13-45"}).status_code == 422


def test_substabelecimento_com_um_substabelecente_e_uma_substabelecida(client):
    payload = {**SUBSTABELECIMENTO, "substabelecentes": SUBSTABELECIMENTO["substabelecentes"][:1],
               "substabelecidos": SUBSTABELECIMENTO["substabelecidos"][:1]}
    texto = _texto(client, _id(client, "substabelecimento", payload))
    assert "OAB/MG nº 111.111, substabelece, SEM RESERVA DE PODERES, à advogada Dra. Beltrana" in texto
    assert "substabelecem" not in texto and "todos com" not in texto
    assert "com endereço profissional à Rua Parceira" in texto


def test_substabelecimento_com_um_substabelecido_homem(client):
    payload = {**SUBSTABELECIMENTO, "substabelecidos": SUBSTABELECIMENTO["substabelecidos"][1:]}
    texto = _texto(client, _id(client, "substabelecimento", payload))
    assert "substabelecem, SEM RESERVA DE PODERES, ao advogado Dr. Ciclano" in texto


def test_substabelecimento_com_varios_continua_no_plural(client):
    texto = _texto(client, _id(client, "substabelecimento", SUBSTABELECIMENTO))
    assert "aos advogados Dra. Beltrana" in texto and "todos com endereço profissional" in texto


def test_distrato_com_ressalva_acrescenta_a_ressalva_a_quitacao(client):
    texto = _texto(client, _id(client, "distrato", {**DISTRATO, "ressalva_exito": True,
                                                    "substabelecer_e_devolver": False}))
    assert "nesse sentido, " + RESSALVA in texto
    assert texto.index("irretratável quitação") < texto.index(RESSALVA)  # qualifica a quitacao
    assert "art. 23 da Lei n.º 8.906/94" in texto
    assert "plena, geral, rasa" not in texto and "substabelecerá" not in texto


def test_distrato_total_usa_a_quitacao_do_modelo_sem_ressalva(client):
    texto = _texto(client, _id(client, "distrato", DISTRATO))
    assert "mútua, total, completa, irrevogável e irretratável quitação" in texto
    assert "renunciando a qualquer direito ou pretensão nesse sentido." in texto
    assert "resta extinto, sem qualquer ônus, penalidade" in texto
    assert "tolerância de qualquer das partes" in texto and "incorpora todas as tratativas" in texto
    assert "foro da comarca eleita no contrato" in texto


def test_distrato_concorda_contratante_e_contratada_em_todo_o_texto(client):
    texto = _texto(client, _id(client, "distrato", {**DISTRATO, **PAGAMENTO}))
    assert "denominada apenas CONTRATADA" in texto and "denominada apenas CONTRATANTE" in texto  # PJ do ORIGEM
    assert "Contratado" not in texto and "A Contratante" not in texto and "O Contratante" not in texto


PAGAMENTO = {"pagamento_valor": 4000, "pagamento_data": "2026-11-30"}


def test_distrato_com_pagamento_traz_valor_data_e_forma(client):
    texto = _texto(client, _id(client, "distrato", {**DISTRATO, **PAGAMENTO, "substabelecer_e_devolver": False}))
    assert ("2.1. A CONTRATANTE pagará à CONTRATADA, até o dia 30/11/2026, o valor de R$ 4.000,00 "
            "(quatro mil reais), por meio de Pix (CNPJ 25.463.159/0001-73).") in texto
    assert texto.index("Por motivos supervenientes") < texto.index("2.1. A CONTRATANTE pagará")
    assert texto.index("2.1. A CONTRATANTE pagará") < texto.index("mútua, total")


def test_distrato_pagamento_com_forma_informada_e_substabelecimento_numerados(client):
    forma = {**PAGAMENTO, "pagamento_forma": "transferência bancária (Banco Exemplo, ag. 0001, c/c 12345-6)."}
    texto = _texto(client, _id(client, "distrato", {**DISTRATO, **forma}))
    assert "por meio de transferência bancária (Banco Exemplo, ag. 0001, c/c 12345-6)." in texto
    assert "Pix (CNPJ" not in texto
    assert "2.3. A CONTRATADA substabelecerá, sem reserva de poderes" in texto


def test_distrato_sem_valor_nao_tem_clausula_de_pagamento(client):
    texto = _texto(client, _id(client, "distrato", {**DISTRATO, "pagamento_data": "2026-11-30"}))
    assert "pagará à CONTRATADA" not in texto and "Pix" not in texto
    assert "2.1. As Partes declaram que os honorários contratuais foram integralmente quitados" in texto
    assert "2.2. A CONTRATADA substabelecerá" in texto


def test_distrato_sem_valor_declara_honorarios_quitados(client):
    texto = _texto(client, _id(client, "distrato", DISTRATO))
    assert ("As Partes declaram que os honorários contratuais foram integralmente quitados, "
            "não havendo valores pendentes.") in texto
    assert "multa" not in texto and "título executivo" not in texto


def test_distrato_com_valor_traz_multa_juros_correcao_e_titulo_executivo(client):
    texto = _texto(client, _id(client, "distrato", {**DISTRATO, **PAGAMENTO}))
    assert "2.2. Em caso de atraso no pagamento, incidirão multa de 10% (dez por cento), juros de 1% (um por cento) ao mês e correção monetária" in texto
    assert "título executivo extrajudicial, nos termos do art. 784, III, do Código de Processo Civil" in texto
    assert "integralmente quitados" not in texto


def test_distrato_substabelecimento_prazo_de_dez_dias_e_recibo(client):
    texto = _texto(client, _id(client, "distrato", DISTRATO))
    assert "em até 10 (dez) dias contados da assinatura deste Distrato" in texto
    assert "prazos processuais que vencerem nesse período, nos termos do art. 5º, § 3º, da Lei n.º 8.906/94" in texto
    assert "mediante recibo" in texto


def test_distrato_contencioso_obriga_o_substabelecimento_junto(client):
    # Mesmo com a opcao desligada, "contencioso" liga a clausula e gera o substabelecimento junto.
    r = _gerar(client, "distrato", {**DISTRATO, "contencioso": True, "substabelecer_e_devolver": False})
    assert r.status_code == 200 and len(r.json()["documentos_gerados_junto"]) == 1
    assert "substabelecerá, sem reserva de poderes" in _texto(client, r.json()["contract_id"])


def test_distrato_contencioso_sem_dados_do_substabelecimento_da_422(client):
    sem = {k: v for k, v in DISTRATO.items() if k != "substabelecimento"}
    assert _gerar(client, "distrato", {**sem, "contencioso": True}).status_code == 422


def test_distrato_consultivo_nao_gera_substabelecimento(client):
    r = _gerar(client, "distrato", {**DISTRATO, "contencioso": False, "substabelecer_e_devolver": False})
    assert r.status_code == 200 and not r.json().get("documentos_gerados_junto")
    assert "substabelecerá" not in _texto(client, r.json()["contract_id"])


@pytest.mark.parametrize("extra", [
    {"pagamento_valor": 4000},  # sem data limite
    {"pagamento_valor": 0, "pagamento_data": "2026-11-30"},
    {"pagamento_valor": 0.001, "pagamento_data": "2026-11-30"},
    {"pagamento_valor": -5, "pagamento_data": "2026-11-30"},
    {"pagamento_valor": 4000, "pagamento_data": "30/11/2026"},
])
def test_distrato_pagamento_invalido_da_422(client, extra):
    assert _gerar(client, "distrato", {**DISTRATO, **extra}).status_code == 422


def test_edicao_nao_troca_o_contrato_de_origem(client):
    cid = _id(client, "aditivo", ADITIVO)
    form = client.get(f"/api/contracts/{cid}/form-data").json()["form_data"]
    form["contrato_origem_id"] = "outro-contrato"
    assert client.put(f"/api/contracts/{cid}", json={"form_data": form}).status_code == 200
    atual = client.get(f"/api/contracts/{cid}/form-data").json()["form_data"]
    assert atual["contrato_origem_id"] == ORIGEM
    assert [d["contract_id"] for d in client.get(f"/api/contracts/{ORIGEM}/documentos").json()["documentos"]] == [cid]


# ── Revisao final (Angela, PR #127) ───────────────────────────────────────


def test_aditivo_cita_so_aditivos_assinados_mas_numera_os_rascunhos(client):
    _assinar(_id(client, "aditivo", ADITIVO))
    _id(client, "aditivo", {**ADITIVO, "objeto": "elaboração de políticas internas"})  # rascunho
    texto = _texto(client, _id(client, "aditivo", {**ADITIVO, "objeto": "treinamento de gestores"}))
    assert "3º ADITIVO CONTRATUAL" in texto
    assert "firmaram o 1º Aditivo" in texto
    assert "políticas internas" not in texto and "2º Aditivo" not in texto


@pytest.mark.parametrize("data", ["20261016", "2026-W42-1", "2026-10-16T10:00", " 2026-10-16"])
def test_data_fora_do_formato_aaaa_mm_dd_da_422(client, data):
    assert _gerar(client, "aditivo", {**ADITIVO, "data_documento": data}).status_code == 422


def test_substabelecimento_um_substabelecente_varios_substabelecidos(client):
    payload = {**SUBSTABELECIMENTO, "substabelecentes": SUBSTABELECIMENTO["substabelecentes"][:1]}
    texto = _texto(client, _id(client, "substabelecimento", payload))
    assert "OAB/MG nº 111.111, substabelece, SEM RESERVA DE PODERES, aos advogados" in texto
    assert "substabelecem" not in texto and "todos com endereço profissional" in texto


# ── Previa (nada e' gravado) ──────────────────────────────────────────────


def _total_contratos() -> int:
    db = SessionLocal()
    try:
        return db.query(ContractDB).count()
    finally:
        db.close()


@pytest.mark.parametrize("tipo,payload,trecho", [
    ("aditivo", ADITIVO, "ADITIVO CONTRATUAL AO CONTRATO"),
    ("distrato", DISTRATO, "DISTRATO DE PRESTA"),
    ("substabelecimento", SUBSTABELECIMENTO, "SUBSTABELECIMENTO SEM RESERVA DE PODERES"),
])
def test_previa_devolve_html_sem_gravar(client, tipo, payload, trecho):
    antes = _total_contratos()
    r = client.post(f"/api/documentos/{tipo}/preview", json=payload)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/html")
    assert trecho in r.text
    assert _total_contratos() == antes


def test_previa_valida_como_a_geracao(client):
    assert client.post("/api/documentos/aditivo/preview", json={**ADITIVO, "data_documento": "2026-13-45"}).status_code == 422
    assert client.post("/api/documentos/procuracao/preview", json=ADITIVO).status_code == 404
    _id(client, "distrato", DISTRATO)
    assert client.post("/api/documentos/distrato/preview", json=DISTRATO).status_code == 409
    _origem_de_outro_advogado()
    assert client.post("/api/documentos/aditivo/preview", json=ADITIVO).status_code == 403


def test_distrato_ressalva_e_o_padrao_do_modelo():
    from app.models.documentos_complementares import DistratoRequest
    d = {k: v for k, v in DISTRATO.items() if k != "ressalva_exito"}
    assert DistratoRequest(**d).ressalva_exito is True


# ── Substabelecimento: formulario simplificado (22/10/2026) ────────────────


def _neutro(**extra) -> dict:
    """Substabelecido como o formulario atual envia: so nome e OAB."""
    return {**SUBSTABELECIMENTO, "substabelecidos": [{"nome": "Beltrana de Tal", "oab": "OAB/MG sob o n.º 3.333"}], **extra}


def test_substabelecido_sem_genero_nem_estado_civil_sai_neutro(client):
    texto = _texto(client, _id(client, "substabelecimento", _neutro()))
    assert "ao(à) advogado(a) Beltrana de Tal, inscrito(a) na OAB/MG sob o n.º 3.333" in texto
    assert "Dra." not in texto and "brasileir" not in texto


def test_substabelecimento_sem_endereco_omite_a_frase(client):
    texto = _texto(client, _id(client, "substabelecimento", _neutro(endereco_substabelecidos="")))
    assert "endereço profissional" not in texto
    assert "OAB/MG sob o n.º 3.333, os poderes que lhes foram conferidos por" in texto


def test_substabelecimento_com_endereco_mantem_a_frase(client):
    texto = _texto(client, _id(client, "substabelecimento", _neutro()))
    assert "com endereço profissional à Rua Parceira, 10, Belo Horizonte/MG, CEP 30.000-000, os poderes" in texto


# ── Distrato que gera tambem o substabelecimento ───────────────────────────


def _derivados_do_origem(tipo: str) -> list:
    db = SessionLocal()
    try:
        return db.query(ContractDB).filter_by(contrato_origem_id=ORIGEM, tipo_contrato=tipo).all()
    finally:
        db.close()


def test_distrato_com_a_opcao_gera_os_dois_documentos(client):
    r = _gerar(client, "distrato", DISTRATO)
    assert r.status_code == 200, r.text
    extras = r.json()["documentos_gerados_junto"]
    assert len(extras) == 1
    # os dois ficam ligados ao contrato de origem
    assert len(_derivados_do_origem("distrato")) == 1
    subs = _derivados_do_origem("substabelecimento")
    assert [x.contract_id for x in subs] == extras
    texto = _texto(client, extras[0])
    assert "SUBSTABELECIMENTO SEM RESERVA DE PODERES" in texto
    assert "Novo Patrono, inscrito(a) na OAB/SP sob o n.º 55.555" in texto
    assert "SÓCIA CF, OAB/MG n.º 111.111" in texto
    assert "Processo n.º 1234567-49.2020.8.13.0024" in texto
    # outorgante e data vem do contrato e do distrato
    assert "EMPRESA EXEMPLO LTDA" in texto and "07 de outubro de 2026" in texto


def test_distrato_sem_a_opcao_gera_so_o_distrato(client):
    payload = {**DISTRATO, "substabelecer_e_devolver": False}
    r = _gerar(client, "distrato", payload)
    assert r.status_code == 200, r.text
    assert r.json()["documentos_gerados_junto"] == []
    assert len(_derivados_do_origem("distrato")) == 1
    assert _derivados_do_origem("substabelecimento") == []


def test_distrato_com_a_opcao_sem_dados_do_substabelecimento_e_rejeitado(client):
    payload = {k: v for k, v in DISTRATO.items() if k != "substabelecimento"}
    assert _gerar(client, "distrato", payload).status_code == 422
    assert _derivados_do_origem("distrato") == []


def test_substabelecimento_do_distrato_aparece_na_lista_do_contrato(client):
    _gerar(client, "distrato", DISTRATO)
    r = client.get(f"/api/contracts/{ORIGEM}/documentos")
    assert r.status_code == 200, r.text
    assert sorted(d["tipo_contrato"] for d in r.json()["documentos"]) == ["distrato", "substabelecimento"]
