"""Numero CNJ -> tribunal; consulta de partes no Legal One (so leitura); importacao por planilha."""
import io

import pytest
from openpyxl import Workbook

from app.services import cnj, legalone_leitura
from tests.test_documentos_complementares import PROC_A, PROC_B, SUBSTABELECIMENTO, _gerar, auth_e_origem  # noqa: F401


def _com_dv(seq: str, ano: str, j: str, tr: str, origem: str) -> str:
    dv = 98 - (int(seq + ano + j + tr + origem + "00") % 97)
    return f"{seq}-{dv:02d}.{ano}.{j}.{tr}.{origem}"


@pytest.mark.parametrize("j,tr,esperado", [
    ("8", "13", "Tribunal de Justiça (MG)"),
    ("8", "26", "Tribunal de Justiça (SP)"),
    ("5", "03", "Tribunal Regional do Trabalho da 3ª Região"),
    ("4", "06", "Tribunal Regional Federal da 6ª Região"),
    ("6", "19", "Tribunal Regional Eleitoral (RJ)"),
    ("3", "00", "Superior Tribunal de Justiça"),
])
def test_tribunal_vem_dos_campos_j_e_tr(j, tr, esperado):
    c = cnj.decodificar(_com_dv("1234567", "2022", j, tr, "0024"))
    assert c and c.tribunal == esperado and c.ano == 2022


def test_aceita_com_ou_sem_mascara_e_devolve_mascarado():
    com = _com_dv("1234567", "2020", "8", "13", "0024")
    sem = com.replace("-", "").replace(".", "")
    assert cnj.normalizar(sem) == com == cnj.normalizar(f" {com} ")
    assert cnj.decodificar(sem).numero == com


@pytest.mark.parametrize("ruim", ["", "123", "1234567-00.2020.8.13.0024", "abc", "1234567-49.2020.0.13.0024"])
def test_digito_ou_formato_errado_nao_decodifica(ruim):
    assert cnj.decodificar(ruim) is None


def test_tr_que_a_tabela_nao_conhece_valida_mas_sem_tribunal():
    c = cnj.decodificar(_com_dv("1234567", "2022", "8", "99", "0024"))
    assert c and c.tribunal is None and c.justica == "Justiça Estadual"


def test_processo_do_substabelecimento_aceita_numero_sem_mascara(client):
    limpo = PROC_A["numero"].replace("-", "").replace(".", "")
    r = _gerar(client, "substabelecimento", {**SUBSTABELECIMENTO, "processos": [{**PROC_A, "numero": limpo}]})
    assert r.status_code == 200, r.text


# --- Legal One: so leitura, silencioso sem credencial -------------------------------------------

def test_sem_credencial_nao_consulta_nada(monkeypatch):
    monkeypatch.delenv("LEGALONE_API_CLIENT_ID", raising=False)
    monkeypatch.delenv("LEGALONE_API_CLIENT_SECRET", raising=False)
    chamadas = []
    monkeypatch.setattr(legalone_leitura.httpx, "get", lambda *a, **k: chamadas.append(1))
    monkeypatch.setattr(legalone_leitura.httpx, "post", lambda *a, **k: chamadas.append(1))
    assert legalone_leitura.partes_do_processo(PROC_A["numero"]) is None
    assert not chamadas


def test_numero_nao_mascarado_nao_entra_no_filtro(monkeypatch):
    monkeypatch.setenv("LEGALONE_API_CLIENT_ID", "x")
    monkeypatch.setenv("LEGALONE_API_CLIENT_SECRET", "y")
    assert legalone_leitura.partes_do_processo("1' or '1'='1") is None


class _Resp:
    def __init__(self, corpo):
        self._corpo = corpo

    def raise_for_status(self):
        pass

    def json(self):
        return self._corpo


def _fake_legalone(monkeypatch, participantes):
    monkeypatch.setenv("LEGALONE_API_CLIENT_ID", "x")
    monkeypatch.setenv("LEGALONE_API_CLIENT_SECRET", "y")
    monkeypatch.setattr(legalone_leitura, "_token", "t")
    monkeypatch.setattr(legalone_leitura, "_token_expira", 9e12)
    pedidos = []

    def get(url, **k):
        pedidos.append((url, k["params"]))
        return _Resp({"value": [{"participants": participantes}] if participantes is not None else []})

    monkeypatch.setattr(legalone_leitura.httpx, "get", get)
    monkeypatch.setattr(legalone_leitura.httpx, "post", lambda *a, **k: pytest.fail("Legal One: so GET"))
    return pedidos


PARTICIPANTES = [
    {"type": "Customer", "contactName": "Cliente A"},
    {"type": "Customer", "contactName": "Cliente B"},
    {"type": "OtherParty", "contactName": "Parte Contrária"},
    {"type": "Party", "contactName": "Outro Participante"},
    {"type": "PersonInCharge", "contactName": "Advogado Responsável"},
]


def test_partes_do_legalone_so_cliente_e_parte_contraria(client, monkeypatch):
    pedidos = _fake_legalone(monkeypatch, PARTICIPANTES)
    r = client.get("/api/processos/cnj", params={"numero": PROC_A["numero"].replace("-", "").replace(".", "")})
    assert r.status_code == 200
    j = r.json()
    assert j["valido"] and j["numero"] == PROC_A["numero"] and j["tribunal"] == "Tribunal de Justiça (MG)"
    assert j["legalone"]["cliente"] == ["Cliente A", "Cliente B"] and j["legalone"]["contraria"] == ["Parte Contrária"]
    assert j["legalone"]["descricao"] == "Cliente A e Cliente B x Parte Contrária"
    assert pedidos[0][0].endswith("/Lawsuits") and PROC_A["numero"] in pedidos[0][1]["$filter"]


def test_processo_fora_do_legalone_devolve_so_o_tribunal(client, monkeypatch):
    _fake_legalone(monkeypatch, None)
    j = client.get("/api/processos/cnj", params={"numero": PROC_B["numero"]}).json()
    assert j["valido"] and j["legalone"] is None and j["tribunal"].startswith("Tribunal Regional Federal")


def test_erro_do_legalone_nao_derruba_a_consulta(client, monkeypatch):
    _fake_legalone(monkeypatch, PARTICIPANTES)

    def quebra(*a, **k):
        raise RuntimeError("segredo: https://x?client_secret=abc")

    monkeypatch.setattr(legalone_leitura.httpx, "get", quebra)
    r = client.get("/api/processos/cnj", params={"numero": PROC_A["numero"]})
    assert r.status_code == 200 and r.json()["legalone"] is None and "segredo" not in r.text


def test_numero_invalido_nao_chama_o_legalone(client, monkeypatch):
    pedidos = _fake_legalone(monkeypatch, PARTICIPANTES)
    j = client.get("/api/processos/cnj", params={"numero": "1234567-00.2020.8.13.0024"}).json()
    assert not j["valido"] and "erro" in j and not pedidos


# --- Planilha --------------------------------------------------------------------------------

def _xlsx(linhas) -> bytes:
    wb = Workbook()
    for l in linhas:
        wb.active.append(l)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _enviar(client, nome, conteudo):
    return client.post("/api/processos/planilha", files={"arquivo": (nome, conteudo)})


def test_planilha_xlsx_com_cabecalho_em_qualquer_ordem(client):
    sem = PROC_B["numero"].replace("-", "").replace(".", "")
    r = _enviar(client, "p.xlsx", _xlsx([
        ["Descrição", "Tipo", "Número do processo"],
        ["A x B", "Agravo", PROC_A["numero"]],
        ["", "", sem],
        ["", "", "1234567-00.2020.8.13.0024"],
        ["", "", PROC_A["numero"]],
    ]))
    assert r.status_code == 200, r.text
    j = r.json()
    p = j["processos"]
    assert (j["validos"], j["invalidos"]) == (2, 2)
    assert p[0]["numero"] == PROC_A["numero"] and p[0]["tipo"] == "Agravo" and p[0]["descricao"] == "A x B"
    assert p[1]["numero"] == PROC_B["numero"] and p[1]["tribunal"].startswith("Tribunal Regional Federal")
    assert not p[2]["valido"] and p[2]["linha"] == 4
    assert p[3]["erro"] == "Número repetido na planilha"


def test_planilha_csv_sem_cabecalho_com_ponto_e_virgula(client):
    csv = f"{PROC_A['numero']};Processo;Fulano x Beltrano\n{PROC_B['numero']};;\n".encode()
    j = _enviar(client, "p.csv", csv).json()
    assert j["validos"] == 2 and j["processos"][0]["descricao"] == "Fulano x Beltrano"


@pytest.mark.parametrize("nome,conteudo", [("p.pdf", b"x"), ("p.xlsx", b"nao e zip"), ("p.csv", b"")])
def test_planilha_ruim_da_422_sem_vazar(client, nome, conteudo):
    r = _enviar(client, nome, conteudo)
    assert r.status_code == 422 and "Traceback" not in r.text


# --- Ajustes da revisao ----------------------------------------------------------------------

def test_planilha_que_descomprime_demais_e_recusada_antes_de_abrir(client, monkeypatch):
    from app.routers import processos

    monkeypatch.setattr(processos, "LIMITE_DESCOMPRIMIDO", 100)  # o xlsx de teste ja passa disso
    abriu = []
    monkeypatch.setattr("openpyxl.load_workbook", lambda *a, **k: abriu.append(1))
    r = _enviar(client, "bomba.xlsx", _xlsx([[PROC_A["numero"]]]))
    assert r.status_code == 422 and "descompactada" in r.text and not abriu


def test_leitura_da_planilha_roda_fora_do_event_loop():
    import inspect

    from app.routers import processos

    assert not inspect.iscoroutinefunction(processos.importar_planilha)  # def: o FastAPI usa o threadpool


def test_numero_invalido_longo_e_truncado_em_60(client):
    j = _enviar(client, "p.csv", ("9" * 300 + "\n").encode()).json()
    assert len(j["processos"][0]["numero"]) == 60


def test_planilha_com_varias_abas_avisa_que_so_a_primeira_foi_lida(client):
    wb = Workbook()
    wb.active.append([PROC_A["numero"]])
    wb.create_sheet("Outra").append([PROC_B["numero"]])
    buf = io.BytesIO()
    wb.save(buf)
    j = _enviar(client, "p.xlsx", buf.getvalue()).json()
    assert "2 abas" in j["aviso"] and j["validos"] == 1
    assert _enviar(client, "p.xlsx", _xlsx([[PROC_A["numero"]]])).json()["aviso"] == ""


@pytest.mark.parametrize("j,tr,esperado", [
    ("5", "00", "Tribunal Superior do Trabalho"), ("6", "00", "Tribunal Superior Eleitoral"),
    ("7", "00", "Superior Tribunal Militar"), ("4", "90", "Conselho da Justiça Federal"),
    ("5", "90", "Conselho Superior da Justiça do Trabalho"),
])
def test_orgaos_nacionais(j, tr, esperado):
    assert cnj.decodificar(_com_dv("1234567", "2022", j, tr, "0000")).tribunal == esperado
